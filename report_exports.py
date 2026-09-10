from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from itertools import groupby
from pathlib import Path
from typing import Any


BRL_FORMAT = 'R$ #,##0.00'
DATE_FORMAT = 'dd/mm/yyyy'
DATETIME_FORMAT = 'dd/mm/yyyy hh:mm'


def _value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    return value if value not in (None, "") else "-"


def generate_xlsx(report: dict, path: Path) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError("A exportação XLSX requer openpyxl. Execute: pip install -r requirements.txt") from exc

    workbook = Workbook()
    workbook.remove(workbook.active)
    navy = "12304A"
    blue = "0757A8"
    pale_blue = "EAF3FB"
    gold = "D99A1B"
    white = "FFFFFF"
    line = Side(style="thin", color="D9E2E7")

    summary = workbook.create_sheet("Resumo")
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:F1")
    summary["A1"] = report["title"]
    summary["A1"].font = Font(size=20, bold=True, color=white)
    summary["A1"].fill = PatternFill("solid", fgColor=navy)
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 34
    summary.merge_cells("A2:F2")
    summary["A2"] = report.get("subtitle", "")
    summary["A2"].font = Font(size=10, color="526575")
    summary["A2"].alignment = Alignment(wrap_text=True)
    row_index = 4
    for label, value in report.get("summary", []):
        summary.cell(row_index, 1, label)
        summary.cell(row_index, 2, _value(value))
        summary.cell(row_index, 1).font = Font(bold=True, color=navy)
        summary.cell(row_index, 1).fill = PatternFill("solid", fgColor=pale_blue)
        summary.cell(row_index, 1).border = Border(bottom=line)
        summary.cell(row_index, 2).border = Border(bottom=line)
        summary.cell(row_index, 2).alignment = Alignment(wrap_text=True, vertical="top")
        row_index += 1
    summary.column_dimensions["A"].width = 27
    summary.column_dimensions["B"].width = 62
    summary.freeze_panes = "A4"

    for section_index, section in enumerate(report.get("sections", []), 1):
        title = str(section["title"])[:31]
        base_title = title
        suffix = 2
        while title in workbook.sheetnames:
            title = f"{base_title[:27]} {suffix}"
            suffix += 1
        sheet = workbook.create_sheet(title)
        sheet.sheet_view.showGridLines = False
        columns = section["columns"]
        rows = section.get("rows", [])
        last_column = get_column_letter(max(1, len(columns)))
        sheet.merge_cells(f"A1:{last_column}1")
        sheet["A1"] = section["title"]
        sheet["A1"].font = Font(size=16, bold=True, color=white)
        sheet["A1"].fill = PatternFill("solid", fgColor=navy)
        sheet["A1"].alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 30
        sheet.merge_cells(f"A2:{last_column}2")
        sheet["A2"] = report.get("subtitle", "")
        sheet["A2"].font = Font(size=10, color="526575")
        sheet["A2"].alignment = Alignment(vertical="center", wrap_text=True)
        header_row = 4
        for col_index, column in enumerate(columns, 1):
            cell = sheet.cell(header_row, col_index, column["label"])
            cell.font = Font(bold=True, color=white)
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=line, right=line, top=line, bottom=line)
        sheet.row_dimensions[header_row].height = 38

        previous_group = None
        group_fill = False
        group_keys = section.get("group_by", [])
        for row_index, record in enumerate(rows, header_row + 1):
            current_group = tuple(record.get(key) for key in group_keys) if group_keys else None
            new_group = bool(group_keys) and current_group != previous_group
            if new_group:
                group_fill = not group_fill
                previous_group = current_group
            for col_index, column in enumerate(columns, 1):
                cell = sheet.cell(row_index, col_index, _value(record.get(column["key"])))
                cell.alignment = Alignment(vertical="top", wrap_text=column.get("wrap", False))
                cell.border = Border(top=Side(style="medium", color=blue) if new_group else Side(), bottom=line)
                if group_keys:
                    cell.fill = PatternFill("solid", fgColor="EEF5FA" if group_fill else "FFFFFF")
                elif row_index % 2 == 1:
                    cell.fill = PatternFill("solid", fgColor="F7FAFC")
                kind = column.get("kind")
                if kind == "currency" and isinstance(cell.value, (int, float)):
                    cell.number_format = BRL_FORMAT
                elif kind == "number" and isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.##'
                elif kind == "date" and isinstance(cell.value, (date, datetime)):
                    cell.number_format = DATE_FORMAT
                elif kind == "datetime" and isinstance(cell.value, (date, datetime)):
                    cell.number_format = DATETIME_FORMAT

            if section.get("highlight_deadline"):
                status = str(record.get("situacao_prazo") or "")
                for key in ("prazo_entrega", "situacao_prazo"):
                    column_index = next((index for index, column in enumerate(columns, 1) if column["key"] == key), None)
                    if not column_index:
                        continue
                    cell = sheet.cell(row_index, column_index)
                    if status.startswith("Atrasada"):
                        cell.fill = PatternFill("solid", fgColor="FDE8E7")
                        cell.font = Font(bold=True, color="B42318")
                    elif status.startswith("Vence"):
                        cell.fill = PatternFill("solid", fgColor="FFF4D6")
                        cell.font = Font(bold=True, color="8A5A00")

        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = f"A{header_row}:{last_column}{max(header_row, len(rows) + header_row)}"
        sheet.auto_filter.add_sort_condition(f"A{header_row + 1}:A{max(header_row + 1, len(rows) + header_row)}")
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = f"1:{header_row}"
        for col_index, column in enumerate(columns, 1):
            samples = [str(column["label"])] + [str(_value(record.get(column["key"]))) for record in rows[:200]]
            longest = max((len(sample) for sample in samples), default=10)
            preferred = column.get("width", 0)
            sheet.column_dimensions[get_column_letter(col_index)].width = min(48, max(preferred, min(longest + 2, 34)))

        if rows:
            table_name = f"LicitaUM{section_index}"
            table = Table(displayName=table_name, ref=f"A{header_row}:{last_column}{len(rows) + header_row}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
            sheet.add_table(table)

    workbook.save(path)
    return path


def _pdf_text(value: Any) -> str:
    from xml.sax.saxutils import escape

    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, (float, Decimal)):
        return f"{float(value):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return escape(str(_value(value))).replace("\n", "<br/>")


def _column_groups(columns: list[dict], maximum: int = 8) -> list[list[dict]]:
    if len(columns) <= maximum:
        return [columns]
    anchors = columns[:2]
    remaining = columns[2:]
    chunk_size = maximum - len(anchors)
    return [anchors + remaining[index:index + chunk_size] for index in range(0, len(remaining), chunk_size)]


def generate_pdf(report: dict, path: Path) -> Path:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A3, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import CondPageBreak, LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise RuntimeError("A exportação PDF requer reportlab. Execute: pip install -r requirements.txt") from exc

    page_size = landscape(A3)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("LicitaUMTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=colors.HexColor("#12304A"), alignment=TA_LEFT, spaceAfter=3 * mm)
    subtitle_style = ParagraphStyle("LicitaUMSubtitle", parent=styles["Normal"], fontSize=8.5, leading=11, textColor=colors.HexColor("#526575"), spaceAfter=5 * mm)
    section_style = ParagraphStyle("LicitaUMSection", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.HexColor("#0757A8"), spaceBefore=3 * mm, spaceAfter=2 * mm)
    cell_style = ParagraphStyle("LicitaUMCell", parent=styles["Normal"], fontName="Helvetica", fontSize=6.2, leading=7.6, textColor=colors.HexColor("#10202A"))
    header_style = ParagraphStyle("LicitaUMHeader", parent=cell_style, fontName="Helvetica-Bold", textColor=colors.white, alignment=TA_LEFT)

    def footer(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D9E2E7"))
        canvas.line(12 * mm, 9 * mm, page_size[0] - 12 * mm, 9 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#667787"))
        canvas.drawString(12 * mm, 5.5 * mm, "LicitaUM - Relatório gerado automaticamente")
        canvas.drawRightString(page_size[0] - 12 * mm, 5.5 * mm, f"Página {document.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(str(path), pagesize=page_size, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=11 * mm, bottomMargin=13 * mm, title=report["title"], author="LicitaUM")
    story = [Paragraph(report["title"], title_style), Paragraph(report.get("subtitle", ""), subtitle_style)]

    summary_rows = []
    for label, value in report.get("summary", []):
        summary_rows.append([Paragraph(f"<b>{_pdf_text(label)}</b>", cell_style), Paragraph(_pdf_text(value), cell_style)])
    if summary_rows:
        summary_table = Table(summary_rows, colWidths=[48 * mm, page_size[0] - 24 * mm - 48 * mm], hAlign="LEFT")
        summary_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF3FB")),
            ("BOX", (0, 0), (-1, -1), .4, colors.HexColor("#D9E2E7")),
            ("INNERGRID", (0, 0), (-1, -1), .25, colors.HexColor("#D9E2E7")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.extend([summary_table, Spacer(1, 5 * mm)])

    available_width = page_size[0] - 24 * mm
    pdf_sections = [section for section in report.get("sections", []) if section.get("rows")]
    for section_index, section in enumerate(pdf_sections):
        if section_index:
            story.append(PageBreak())
        story.append(Paragraph(section["title"], section_style))
        rows = section.get("rows", [])
        pdf_columns = section.get("pdf_columns")
        column_groups = [pdf_columns] if pdf_columns else _column_groups(section["columns"], maximum=12)
        group_keys = section.get("group_by", [])
        grouped_records = []
        if group_keys and pdf_columns:
            key_function = lambda record: tuple(record.get(key) for key in group_keys)
            grouped_records = [(key, list(records)) for key, records in groupby(rows, key=key_function)]

        report_groups = grouped_records if grouped_records else [(None, rows)]
        for record_group_index, (_, records) in enumerate(report_groups):
            if record_group_index:
                story.append(Spacer(1, 4 * mm))
            for column_group_index, columns in enumerate(column_groups, 1):
                if len(column_groups) > 1:
                    story.append(Paragraph(f"{section['title']} - informações {column_group_index} de {len(column_groups)}", subtitle_style))
                if grouped_records:
                    story.append(CondPageBreak(30 * mm))
                width_units = [max(1, column.get("pdf_weight", 1)) for column in columns]
                unit_total = sum(width_units)
                col_widths = [available_width * unit / unit_total for unit in width_units]
                table_rows = []
                header_row = 0
                data_start = 1
                if grouped_records:
                    context = " &nbsp;&nbsp; | &nbsp;&nbsp; ".join(
                        f"<b>{_pdf_text(label)}:</b> {_pdf_text(records[0].get(key))}"
                        for label, key in section.get("group_context", [])
                    )
                    table_rows.append([Paragraph(context, header_style)] + [""] * (len(columns) - 1))
                    header_row = 1
                    data_start = 2
                table_rows.append([Paragraph(_pdf_text(column["label"]), header_style) for column in columns])
                for record in records:
                    values = []
                    for column in columns:
                        value = record.get(column["key"])
                        if column.get("kind") == "currency" and isinstance(value, (int, float, Decimal)):
                            value = "R$ " + _pdf_text(value)
                        values.append(Paragraph(_pdf_text(value), cell_style))
                    table_rows.append(values)
                table = LongTable(table_rows, colWidths=col_widths, repeatRows=data_start, hAlign="LEFT")
                table_commands = [
                    ("BACKGROUND", (0, header_row), (-1, header_row), colors.HexColor("#0757A8")),
                    ("ROWBACKGROUNDS", (0, data_start), (-1, -1), [colors.white, colors.HexColor("#F4F8FA")]),
                    ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#C9D5DB")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
                if grouped_records:
                    table_commands.extend([
                        ("SPAN", (0, 0), (-1, 0)),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12304A")),
                        ("TOPPADDING", (0, 0), (-1, 0), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, 0), 5),
                    ])
                if section.get("highlight_deadline"):
                    deadline_index = next((index for index, column in enumerate(columns) if column["key"] == "prazo_entrega"), None)
                    status_index = next((index for index, column in enumerate(columns) if column["key"] == "situacao_prazo"), None)
                    for index, record in enumerate(records, data_start):
                        status = str(record.get("situacao_prazo") or "")
                        if status.startswith("Atrasada"):
                            for column_index in (deadline_index, status_index):
                                if column_index is not None:
                                    table_commands.extend([
                                        ("BACKGROUND", (column_index, index), (column_index, index), colors.HexColor("#FDE8E7")),
                                        ("TEXTCOLOR", (column_index, index), (column_index, index), colors.HexColor("#B42318")),
                                    ])
                        elif status.startswith("Vence"):
                            for column_index in (deadline_index, status_index):
                                if column_index is not None:
                                    table_commands.append(("BACKGROUND", (column_index, index), (column_index, index), colors.HexColor("#FFF4D6")))
                table.setStyle(TableStyle(table_commands))
                story.append(table)
                if column_group_index < len(column_groups):
                    story.extend([PageBreak(), Paragraph(section["title"], section_style)])

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return path
