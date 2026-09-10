from __future__ import annotations

import json
import hashlib
import os
import re
import shutil
import sqlite3
import subprocess
import base64
import mimetypes
import uuid
from contextlib import contextmanager
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
DB_PATH = ROOT / "licitacoes.db"
STATIC = ROOT / "static"
OUTPUTS = ROOT / "propostas_geradas"
UPLOADS = ROOT / "uploads"
EXPORT_TEMP = ROOT / "data" / "exportacoes_temporarias"
TEMPLATE_SOURCE = Path(r"C:\Users\selector\Downloads\Proposta Vogen Cosmetics (10).docx")
TEMPLATE_COPY = ROOT / "modelo_proposta_vogen.docx"


def current_business_date() -> date:
    return datetime.now(ZoneInfo("America/Sao_Paulo")).date()


STATUS = [
    "Em cadastro de preços",
    "Futura licitação",
    "Licitação passada",
    "Proposta enviada",
    "Julgado e Habilitado",
    "Aguardando Habilitação",
    "Adjudicada",
    "Nota de empenho emitida",
    "Entregue",
    "Perdido",
]


@contextmanager
def connect():
    con = sqlite3.connect(DB_PATH, timeout=15)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys = ON")
    con.execute("PRAGMA busy_timeout = 10000")
    try:
        yield con
        con.commit()
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()


def money(value) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    cleaned = str(value).replace("R$", "").replace(" ", "").strip()
    has_comma = "," in cleaned
    has_dot = "." in cleaned
    if has_comma:
        cleaned = cleaned.replace(".", "").replace(",", ".")
    elif has_dot:
        parts = cleaned.split(".")
        if len(parts[-1]) in (1, 2):
            cleaned = "".join(parts[:-1]) + "." + parts[-1]
        else:
            cleaned = cleaned.replace(".", "")
    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return Decimal("0")


def brl(value) -> str:
    amount = money(value).quantize(Decimal("0.01"))
    whole, cents = f"{amount:.2f}".split(".")
    groups = []
    while whole:
        groups.insert(0, whole[-3:])
        whole = whole[:-3]
    return f"R$ {'.'.join(groups)},{cents}"


def parse_date(value: str | None) -> date | None:
    if not value:
        return None
    text = str(value).replace("Z", "")
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        try:
            return datetime.strptime(text, "%Y-%m-%d").date()
        except ValueError:
            return None


def safe_name(text: str) -> str:
    allowed = "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789-_ "
    return "".join(ch for ch in text if ch in allowed).strip().replace(" ", "_")[:90] or "proposta"


def proposal_filename(tender: dict, items: list[dict]) -> str:
    item_numbers = [str(item.get("item") or item.get("id")) for item in items]
    label = "Item" if len(item_numbers) == 1 else "Itens"
    items_text = " ".join(item_numbers)
    return safe_name(f"Proposta {label} {items_text} Pregao {tender['pregão']} UASG {tender['uasg']}") + ".docx"


def set_paragraph_text(paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def set_cell_text(cell, text: str) -> None:
    cell.text = text


def row_to_dict(row: sqlite3.Row) -> dict:
    return dict(row)


def hash_password(password: str) -> str:
    if not password:
        return ""
    return hashlib.sha256(password.encode("utf-8")).hexdigest()


def init_db() -> None:
    if TEMPLATE_SOURCE.exists() and not TEMPLATE_COPY.exists():
        TEMPLATE_COPY.write_bytes(TEMPLATE_SOURCE.read_bytes())

    with connect() as con:
        con.execute("PRAGMA journal_mode = WAL")
        con.execute("PRAGMA synchronous = NORMAL")
        con.executescript(
            """
            CREATE TABLE IF NOT EXISTS tenders (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                pregão TEXT NOT NULL,
                uasg TEXT NOT NULL,
                órgão TEXT NOT NULL,
                localidade TEXT,
                plataforma TEXT NOT NULL DEFAULT 'ComprasNet',
                data_limite TEXT,
                data_proposta TEXT,
                status TEXT NOT NULL DEFAULT 'Em cadastro de preços',
                modalidade TEXT NOT NULL DEFAULT 'Pregão Eletrônico',
                observação TEXT,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );

            CREATE TABLE IF NOT EXISTS items (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                tender_id INTEGER NOT NULL REFERENCES tenders(id) ON DELETE CASCADE,
                item TEXT,
                marca TEXT,
                modelo TEXT,
                referência TEXT,
                qtd INTEGER NOT NULL DEFAULT 1,
                valor_unitário REAL NOT NULL DEFAULT 0,
                link_usa TEXT,
                link_br TEXT,
                link_referência TEXT,
                valor_cadastro REAL,
                valor_mínimo REAL,
                responsável_preço TEXT,
                status TEXT NOT NULL DEFAULT 'Em cadastro de preços',
                observação TEXT
            );

            CREATE TABLE IF NOT EXISTS orders (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                item_id INTEGER NOT NULL REFERENCES items(id) ON DELETE CASCADE,
                qtd_empenhada REAL,
                prazo_entrega TEXT,
                ordem_fornecimento TEXT,
                endereço_entrega TEXT,
                nota_empenho TEXT,
                status TEXT NOT NULL DEFAULT 'Pendente',
                observação TEXT
            );
            """
        )
        count = con.execute("SELECT COUNT(*) FROM tenders").fetchone()[0]
        if count:
            return

        samples = [
            (
                "90040/2026",
                "000000",
                "ÓRGÃO A DEFINIR",
                "",
                "2026-07-15T09:00",
                "Futura licitação",
                [
                    ("", "", "", "Itens ainda em análise", 1, 0, "", "", "", None, None, ""),
                ],
            ),
            (
                "90012/2025",
                "930105",
                "CÂMARA MUNICIPAL DE JACAREÍ - SP",
                "Jacareí-SP",
                "2025-03-04T09:00",
                "Em cadastro de preços",
                [
                    ("6", "Hollyland", "Lark M2 Combo", "Hollyland Lark M2 Combo", 6, 1237.80, "bhphotovideo", "Mercado Livre", "", 2000, 800, ""),
                    ("7", "AKG", "K414P", "AKG K414P", 10, 283.30, "", "Mercado Livre", "", 400, 380, ""),
                    ("10", "Sony", "BP-U70", "Sony BP-U70", 4, 3750.40, "bhphotovideo", "", "", 4200, 3500, ""),
                    ("25", "Hollyland", "Solidcom C1 Pro Hub 8S", "Hollyland Solidcom C1 Pro Hub 8S", 1, 47666.97, "", "", "", 50000, 47000, ""),
                ],
            ),
            (
                "90028/2025",
                "158517",
                "UNIVERSIDADE FEDERAL DA FRONTEIRA SUL",
                "",
                "",
                "Aguardando Habilitação",
                [
                    ("44", "Pioneer", "RMX-1000", "Pioneer RMX-1000", 1, 11499, "", "", "https://www.pioneerdj.com/", None, None, ""),
                ],
            ),
            (
                "90036/2025",
                "158410",
                "INST. FED. DE EDUC. TEC. BAHIA/CAMPUS EUNÁPOLIS",
                "Eunápolis-BA",
                "",
                "Nota de empenho emitida",
                [
                    ("42", "Yamaha", "TF5", "Yamaha TF5", 1, 25969, "", "", "Yamaha", None, None, "99"),
                    ("43", "Epson", "PowerLite E20", "Epson PowerLite E20", 8, 3299, "", "", "Epson", None, None, "129"),
                ],
            ),
        ]
        for pregão, uasg, órgão, localidade, data_limite, status, items in samples:
            cur = con.execute(
                "INSERT INTO tenders (pregão, uasg, órgão, localidade, data_limite, status) VALUES (?, ?, ?, ?, ?, ?)",
                (pregão, uasg, órgão, localidade, data_limite, status),
            )
            tender_id = cur.lastrowid
            for item in items:
                cur_item = con.execute(
                    """
                    INSERT INTO items (
                        tender_id, item, marca, modelo, referência, qtd, valor_unitário,
                        link_usa, link_br, link_referência, valor_cadastro, valor_mínimo,
                        observação
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (tender_id, *item),
                )
                if status == "Nota de empenho emitida":
                    con.execute(
                        """
                        INSERT INTO orders (item_id, prazo_entrega, ordem_fornecimento, endereço_entrega, status)
                        VALUES (?, ?, ?, ?, ?)
                        """,
                        (
                            cur_item.lastrowid,
                            "2026-01-29" if item[0] == "42" else "2026-01-30",
                            item[-1],
                            "Avenida David Jonas Fadini BR 101, Rosa Neto, Eunápolis - BA. CEP 45823-221",
                            "Pendente",
                        ),
                    )


def ensure_schema() -> None:
    with connect() as con:
        tender_cols = {row["name"] for row in con.execute("PRAGMA table_info(tenders)")}
        if "plataforma" not in tender_cols:
            con.execute("ALTER TABLE tenders ADD COLUMN plataforma TEXT NOT NULL DEFAULT 'ComprasNet'")
        if "data_proposta" not in tender_cols:
            con.execute("ALTER TABLE tenders ADD COLUMN data_proposta TEXT")
        item_cols = {row["name"] for row in con.execute("PRAGMA table_info(items)")}
        if "valor_ganho" not in item_cols:
            con.execute("ALTER TABLE items ADD COLUMN valor_ganho REAL")
        if "valor_sigiloso" not in item_cols:
            con.execute("ALTER TABLE items ADD COLUMN valor_sigiloso INTEGER NOT NULL DEFAULT 0")
        if "lote" not in item_cols:
            con.execute("ALTER TABLE items ADD COLUMN lote TEXT")
        if "opção_produto" not in item_cols:
            con.execute("ALTER TABLE items ADD COLUMN opção_produto TEXT")
        if "selecionado_cadastro" not in item_cols:
            con.execute("ALTER TABLE items ADD COLUMN selecionado_cadastro INTEGER NOT NULL DEFAULT 1")
        order_cols = {row["name"] for row in con.execute("PRAGMA table_info(orders)")}
        if "pagamento_recebido" not in order_cols:
            con.execute("ALTER TABLE orders ADD COLUMN pagamento_recebido INTEGER NOT NULL DEFAULT 0")
        if "qtd_empenhada" not in order_cols:
            con.execute("ALTER TABLE orders ADD COLUMN qtd_empenhada REAL")
        if "group_id" not in order_cols:
            con.execute("ALTER TABLE orders ADD COLUMN group_id TEXT")
        if "origem" not in order_cols:
            con.execute("ALTER TABLE orders ADD COLUMN origem TEXT NOT NULL DEFAULT 'Empenho'")
        if "órgão_solicitante" not in order_cols:
            con.execute("ALTER TABLE orders ADD COLUMN órgão_solicitante TEXT")
        con.execute(
            """
            UPDATE orders
            SET qtd_empenhada = (SELECT qtd FROM items WHERE items.id = orders.item_id)
            WHERE qtd_empenhada IS NULL OR qtd_empenhada <= 0
            """
        )
        con.execute(
            "UPDATE orders SET group_id = 'legacy-' || id WHERE group_id IS NULL OR group_id = ''"
        )
        con.execute(
            """
            UPDATE items
            SET status = 'Adjudicada'
            WHERE status = 'Entregue'
              AND EXISTS (SELECT 1 FROM orders WHERE orders.item_id = items.id)
            """
        )
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS attachments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                tender_id INTEGER REFERENCES tenders(id) ON DELETE CASCADE,
                item_id INTEGER REFERENCES items(id) ON DELETE CASCADE,
                tipo TEXT NOT NULL,
                filename TEXT NOT NULL,
                stored_name TEXT NOT NULL,
                content_type TEXT,
                uploaded_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nome TEXT NOT NULL,
                email TEXT NOT NULL UNIQUE,
                perfil TEXT NOT NULL DEFAULT 'Usuário',
                senha_hash TEXT,
                ativo INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        con.executescript(
            """
            CREATE INDEX IF NOT EXISTS idx_tenders_data_limite ON tenders(data_limite);
            CREATE INDEX IF NOT EXISTS idx_tenders_status ON tenders(status);
            CREATE INDEX IF NOT EXISTS idx_items_tender_id ON items(tender_id);
            CREATE INDEX IF NOT EXISTS idx_items_status ON items(status);
            CREATE INDEX IF NOT EXISTS idx_orders_item_id ON orders(item_id);
            CREATE INDEX IF NOT EXISTS idx_orders_group_id ON orders(group_id);
            CREATE INDEX IF NOT EXISTS idx_attachments_tender_id ON attachments(tender_id);
            CREATE INDEX IF NOT EXISTS idx_attachments_item_id ON attachments(item_id);
            """
        )
        users = con.execute("SELECT COUNT(*) FROM users").fetchone()[0]
        if not users:
            con.execute(
                "INSERT INTO users (nome, email, perfil, senha_hash) VALUES (?, ?, ?, ?)",
                ("Administrador", "admin@licitaum.local", "Administrador", hash_password("admin")),
            )


def refresh_past_tenders(con: sqlite3.Connection) -> None:
    now = datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%Y-%m-%dT%H:%M")
    con.execute(
        """
        UPDATE tenders
        SET status = 'Licitação passada'
        WHERE status = 'Futura licitação'
          AND data_limite IS NOT NULL
          AND data_limite != ''
          AND datetime(data_limite) < datetime(?)
        """,
        (now,),
    )


def all_tenders(query: str = "", status: str = "") -> list[dict]:
    with connect() as con:
        refresh_past_tenders(con)
        sql = """
            SELECT t.*,
                   COUNT(i.id) AS itens,
                   COALESCE(SUM(CASE WHEN COALESCE(i.selecionado_cadastro, 1) = 1 THEN 1 ELSE 0 END), 0) AS itens_ativos,
                   COALESCE(SUM(
                       CASE
                           WHEN COALESCE(i.selecionado_cadastro, 1) = 1 AND COALESCE(i.valor_sigiloso, 0) = 1 THEN 1
                           ELSE 0
                       END
                   ), 0) AS itens_sigilosos,
                   COALESCE(SUM(
                       CASE
                           WHEN COALESCE(i.selecionado_cadastro, 1) = 0 THEN 0
                           WHEN COALESCE(i.valor_cadastro, 0) <= 0 THEN 0
                           WHEN COALESCE(i.valor_mínimo, 0) <= 0 THEN 0
                           WHEN COALESCE(i.valor_sigiloso, 0) = 1 THEN 1
                           WHEN COALESCE(i.valor_unitário, 0) > 0 THEN 1
                           ELSE 0
                       END
                   ), 0) AS itens_precificados,
                   COALESCE(SUM(
                       CASE
                           WHEN COALESCE(i.selecionado_cadastro, 1) = 0 THEN 0
                           WHEN COALESCE(i.valor_sigiloso, 0) = 1 THEN 0
                           ELSE i.qtd * i.valor_unitário
                       END
                   ), 0) AS valor_total,
                   COALESCE(SUM(
                       CASE
                           WHEN COALESCE(i.selecionado_cadastro, 1) = 0 THEN 0
                           WHEN COALESCE(i.valor_mínimo, 0) <= 0 THEN 0
                           ELSE i.qtd * i.valor_mínimo
                       END
                   ), 0) AS valor_minimo_total
            FROM tenders t
            LEFT JOIN items i ON i.tender_id = t.id
        """
        where, args = [], []
        if query:
            where.append("(t.pregão LIKE ? OR t.uasg LIKE ? OR t.órgão LIKE ?)")
            like = f"%{query}%"
            args.extend([like, like, like])
        if status:
            where.append("t.status = ?")
            args.append(status)
        if where:
            sql += " WHERE " + " AND ".join(where)
        sql += " GROUP BY t.id ORDER BY t.created_at DESC, t.id DESC"
        return [row_to_dict(row) for row in con.execute(sql, args)]


def effective_unit_value(item: dict) -> Decimal:
    valor_ganho = money(item.get("valor_ganho"))
    return valor_ganho if valor_ganho > 0 else money(item.get("valor_unitário"))


def attachments_for(tender_id: int | None = None, item_id: int | None = None) -> list[dict]:
    with connect() as con:
        where, args = [], []
        if tender_id:
            where.append("tender_id = ?")
            args.append(tender_id)
        if item_id:
            where.append("item_id = ?")
            args.append(item_id)
        sql = "SELECT * FROM attachments"
        if where:
            sql += " WHERE " + " AND ".join(where)
        sql += " ORDER BY uploaded_at DESC, id DESC"
        return [row_to_dict(row) for row in con.execute(sql, args)]


def items_with_orders(con: sqlite3.Connection, tender_id: int | None = None) -> list[dict]:
    where = " WHERE tender_id = ?" if tender_id is not None else ""
    args = (tender_id,) if tender_id is not None else ()
    item_rows = con.execute(
        f"SELECT * FROM items{where} ORDER BY tender_id, CAST(item AS INTEGER), id",
        args,
    ).fetchall()
    items = [row_to_dict(row) for row in item_rows]
    if not items:
        return []

    item_by_id = {int(item["id"]): item for item in items}
    placeholders = ", ".join("?" for _ in item_by_id)
    order_rows = con.execute(
        f"SELECT * FROM orders WHERE item_id IN ({placeholders}) ORDER BY id",
        list(item_by_id),
    ).fetchall()

    for item in items:
        item["orders"] = []

    for row in order_rows:
        order = row_to_dict(row)
        item = item_by_id.get(int(order["item_id"]))
        if not item:
            continue
        quantity = money(order.get("qtd_empenhada"))
        if quantity <= 0:
            quantity = money(item.get("qtd"))
        order["qtd_empenhada"] = float(quantity)
        item["orders"].append(order)

    for item in items:
        committed = sum(
            money(order.get("qtd_empenhada"))
            for order in item["orders"]
            if order.get("origem") != "Carona"
        )
        won = money(item.get("qtd"))
        item["qtd_empenhada_total"] = float(committed)
        item["qtd_pendente"] = float(max(won - committed, Decimal("0")))
    return items


def tender_detail(tender_id: int) -> dict:
    with connect() as con:
        tender = con.execute("SELECT * FROM tenders WHERE id = ?", (tender_id,)).fetchone()
        if not tender:
            raise KeyError("Licitação não encontrada")
        data = row_to_dict(tender)
        data["items"] = items_with_orders(con, tender_id)
        data["attachments"] = attachments_for(tender_id=tender_id)
        data["valor_total"] = sum(
            money(row["qtd"]) * money(row.get("valor_unitário"))
            for row in data["items"]
            if int(row.get("selecionado_cadastro") if row.get("selecionado_cadastro") is not None else 1)
            and not int(row.get("valor_sigiloso") or 0)
        )
        data["valor_minimo_total"] = sum(
            money(row["qtd"]) * money(row.get("valor_mínimo"))
            for row in data["items"]
            if int(row.get("selecionado_cadastro") if row.get("selecionado_cadastro") is not None else 1)
            and money(row.get("valor_mínimo")) > 0
        )
        return data


def application_state() -> list[dict]:
    tenders = all_tenders()
    if not tenders:
        return []

    by_id = {int(tender["id"]): tender for tender in tenders}
    for tender in tenders:
        tender["items"] = []
        tender["attachments"] = []

    with connect() as con:
        items = items_with_orders(con)
        attachments = con.execute(
            "SELECT * FROM attachments ORDER BY uploaded_at DESC, id DESC"
        ).fetchall()

    for item in items:
        tender = by_id.get(int(item["tender_id"]))
        if tender:
            tender["items"].append(item)

    for row in attachments:
        attachment = row_to_dict(row)
        tender_id = attachment.get("tender_id")
        tender = by_id.get(int(tender_id)) if tender_id else None
        if tender:
            tender["attachments"].append(attachment)

    return tenders


def export_datetime(value) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace("Z", ""))
    except ValueError:
        return None


def export_column(key: str, label: str, kind: str = "text", width: int = 0, pdf_weight: float = 1) -> dict:
    return {"key": key, "label": label, "kind": kind, "width": width, "pdf_weight": pdf_weight, "wrap": kind == "text"}


TENDER_EXPORT_COLUMNS = [
    export_column("pregao", "Pregão", width=18),
    export_column("uasg", "UASG", width=13),
    export_column("orgao", "Órgão", width=34, pdf_weight=2.2),
    export_column("localidade", "Localidade", width=20),
    export_column("plataforma", "Plataforma", width=14),
    export_column("modalidade", "Modalidade", width=20),
    export_column("data_limite", "Data e hora", "datetime", width=18),
    export_column("data_proposta", "Data da proposta", "date", width=17),
    export_column("status_licitacao", "Status da licitação", width=23),
]

ITEM_EXPORT_COLUMNS = TENDER_EXPORT_COLUMNS + [
    export_column("lote", "Lote", width=10),
    export_column("item", "Item", width=9),
    export_column("tipo_produto", "Produto", width=13),
    export_column("cadastrar", "Cadastrar/Disputar", width=18),
    export_column("marca", "Marca", width=18),
    export_column("modelo", "Modelo", width=25, pdf_weight=1.6),
    export_column("referencia", "Referência", width=28, pdf_weight=1.8),
    export_column("quantidade", "Qtd. ganha/solicitada", "number", width=18),
    export_column("valor_sigiloso", "Referência sigilosa", width=17),
    export_column("valor_unitario_referencia", "Valor unit. referência", "currency", width=20),
    export_column("valor_total_referencia", "Total referência", "currency", width=18),
    export_column("valor_cadastro", "Valor cadastro", "currency", width=17),
    export_column("valor_total_cadastro", "Total cadastro", "currency", width=17),
    export_column("valor_minimo", "Valor mínimo", "currency", width=16),
    export_column("valor_total_minimo", "Total mínimo", "currency", width=16),
    export_column("valor_ganho", "Valor ganho unit.", "currency", width=18),
    export_column("valor_total_ganho", "Total ganho", "currency", width=17),
    export_column("valor_total_pendente", "Total pendente", "currency", width=17),
    export_column("status_item", "Status do item", width=23),
    export_column("qtd_empenhada", "Qtd. empenhada", "number", width=16),
    export_column("qtd_pendente", "Qtd. pendente", "number", width=15),
    export_column("responsavel_preco", "Responsável pelo preço", width=23),
    export_column("link_fornecedor", "Link fornecedor", width=35, pdf_weight=2),
    export_column("link_br", "Link BR", width=30, pdf_weight=1.8),
    export_column("link_usa", "Link USA", width=30, pdf_weight=1.8),
    export_column("observacao_item", "Observação do item", width=38, pdf_weight=2.4),
]

ORDER_EXPORT_COLUMNS = TENDER_EXPORT_COLUMNS + [
    export_column("grupo", "Grupo da encomenda", width=20),
    export_column("origem", "Origem", width=12),
    export_column("orgao_solicitante", "Órgão solicitante", width=30, pdf_weight=2),
    export_column("lote", "Lote", width=10),
    export_column("item", "Item", width=9),
    export_column("marca", "Marca", width=18),
    export_column("modelo", "Modelo", width=25, pdf_weight=1.6),
    export_column("referencia", "Referência", width=28, pdf_weight=1.8),
    export_column("qtd_empenhada", "Qtd. empenhada", "number", width=16),
    export_column("valor_unitario", "Valor unitário", "currency", width=17),
    export_column("valor_total", "Valor total", "currency", width=17),
    export_column("prazo_entrega", "Prazo de entrega", "date", width=17),
    export_column("ordem_fornecimento", "Ordem de fornecimento", width=21),
    export_column("nota_empenho", "Nota de empenho", width=19),
    export_column("status_encomenda", "Status da encomenda", width=20),
    export_column("pagamento", "Pagamento", width=19),
    export_column("endereco_entrega", "Endereço de entrega", width=48, pdf_weight=3),
    export_column("observacao_encomenda", "Observação da encomenda", width=40, pdf_weight=2.5),
    export_column("status_item", "Status do item", width=23),
    export_column("observacao_item", "Observação do item", width=38, pdf_weight=2.4),
]

DOCUMENT_EXPORT_COLUMNS = [
    export_column("tipo", "Tipo", width=16),
    export_column("item", "Item vinculado", width=15),
    export_column("arquivo", "Arquivo", width=45, pdf_weight=3),
    export_column("content_type", "Formato", width=27),
    export_column("enviado_em", "Enviado em", "datetime", width=20),
]


def tender_export_base(tender: dict) -> dict:
    proposal_date = export_datetime(tender.get("data_proposta"))
    return {
        "pregao": tender.get("pregão") or "-",
        "uasg": tender.get("uasg") or "-",
        "orgao": tender.get("órgão") or "-",
        "localidade": tender.get("localidade") or "-",
        "plataforma": tender.get("plataforma") or "ComprasNet",
        "modalidade": tender.get("modalidade") or "Pregão Eletrônico",
        "data_limite": export_datetime(tender.get("data_limite")),
        "data_proposta": proposal_date.date() if proposal_date else None,
        "status_licitacao": tender.get("status") or "-",
    }


def item_export_record(tender: dict, item: dict) -> dict:
    quantity = money(item.get("qtd"))
    reference = money(item.get("valor_unitário"))
    registration = money(item.get("valor_cadastro"))
    minimum = money(item.get("valor_mínimo"))
    won = money(item.get("valor_ganho"))
    pending_quantity = money(item.get("qtd_pendente"))
    secret = bool(item.get("valor_sigiloso"))
    selected = int(item.get("selecionado_cadastro") if item.get("selecionado_cadastro") is not None else 1) != 0
    option = str(item.get("opção_produto") or "").strip()
    return {
        **tender_export_base(tender),
        "lote": item.get("lote") or "Avulso",
        "item": item.get("item") or "-",
        "tipo_produto": option or "Principal",
        "cadastrar": "Sim" if selected else "Não disputar",
        "marca": item.get("marca") or "-",
        "modelo": item.get("modelo") or "-",
        "referencia": item.get("referência") or "-",
        "quantidade": float(quantity),
        "valor_sigiloso": "Sim" if secret else "Não",
        "valor_unitario_referencia": None if secret else float(reference),
        "valor_total_referencia": None if secret else float(reference * quantity),
        "valor_cadastro": float(registration),
        "valor_total_cadastro": float(registration * quantity),
        "valor_minimo": float(minimum),
        "valor_total_minimo": float(minimum * quantity),
        "valor_ganho": float(won),
        "valor_total_ganho": float(won * quantity),
        "valor_total_pendente": float(won * pending_quantity),
        "status_item": item.get("status") or "-",
        "qtd_empenhada": float(money(item.get("qtd_empenhada_total"))),
        "qtd_pendente": float(pending_quantity),
        "responsavel_preco": item.get("responsável_preço") or "-",
        "link_fornecedor": item.get("link_referência") or "-",
        "link_br": item.get("link_br") or "-",
        "link_usa": item.get("link_usa") or "-",
        "observacao_item": item.get("observação") or "-",
    }


def order_export_record(tender: dict, item: dict, order: dict) -> dict:
    quantity = money(order.get("qtd_empenhada"))
    unit_value = effective_unit_value(item)
    return {
        **tender_export_base(tender),
        "grupo": order.get("group_id") or f"order-{order.get('id')}",
        "origem": order.get("origem") or "Empenho",
        "orgao_solicitante": order.get("órgão_solicitante") or "-",
        "lote": item.get("lote") or "Avulso",
        "item": item.get("item") or "-",
        "marca": item.get("marca") or "-",
        "modelo": item.get("modelo") or "-",
        "referencia": item.get("referência") or "-",
        "qtd_empenhada": float(quantity),
        "valor_unitario": float(unit_value),
        "valor_total": float(quantity * unit_value),
        "prazo_entrega": export_datetime(order.get("prazo_entrega")),
        "ordem_fornecimento": order.get("ordem_fornecimento") or "-",
        "nota_empenho": order.get("nota_empenho") or "-",
        "status_encomenda": order.get("status") or "-",
        "pagamento": "Pago" if order.get("pagamento_recebido") else "Pendente",
        "endereco_entrega": order.get("endereço_entrega") or "-",
        "observacao_encomenda": order.get("observação") or "-",
        "status_item": item.get("status") or "-",
        "observacao_item": item.get("observação") or "-",
        "_order_id": int(order.get("id") or 0),
    }


def all_export_rows() -> tuple[list[dict], list[dict]]:
    item_rows, order_rows = [], []
    for tender in application_state():
        for item in tender.get("items", []):
            item_record = item_export_record(tender, item)
            item_record["_item_id"] = int(item.get("id") or 0)
            item_rows.append(item_record)
            for order in item.get("orders", []):
                order_rows.append(order_export_record(tender, item, order))
    return item_rows, order_rows


def report_total(rows: list[dict], key: str) -> Decimal:
    return sum((money(row.get(key)) for row in rows), Decimal("0"))


def sector_export_report(sector: str, selected_ids: set[int] | None = None) -> dict:
    item_rows, order_rows = all_export_rows()
    sector_names = {
        "won": "Itens Ganhos",
        "proposal": "Itens em Proposta",
        "orders": "Encomendas",
        "finished": "Finalizados",
    }
    if sector not in sector_names:
        raise ValueError("Tipo de exportação inválido.")

    if sector == "won":
        rows = [row for row in item_rows if row["status_item"] in ("Aguardando Habilitação", "Julgado e Habilitado", "Adjudicada") and money(row["qtd_pendente"]) > 0]
        id_key, columns, total_key = "_item_id", ITEM_EXPORT_COLUMNS, "valor_total_pendente"
    elif sector == "proposal":
        rows = [row for row in item_rows if row["status_item"] == "Proposta enviada"]
        id_key, columns, total_key = "_item_id", ITEM_EXPORT_COLUMNS, "valor_total_ganho"
    elif sector == "orders":
        rows = [row for row in order_rows if row["status_encomenda"] != "Entregue"]
        id_key, columns, total_key = "_order_id", ORDER_EXPORT_COLUMNS, "valor_total"
    else:
        rows = [row for row in order_rows if row["status_encomenda"] == "Entregue"]
        id_key, columns, total_key = "_order_id", ORDER_EXPORT_COLUMNS, "valor_total"

    if selected_ids is not None:
        rows = [row for row in rows if row[id_key] in selected_ids]
    title = f"LicitaUM - {sector_names[sector]}"
    unique_tenders = {(row["pregao"], row["uasg"]) for row in rows}
    summary = [
        ("Registros exportados", len(rows)),
        ("Licitações relacionadas", len(unique_tenders)),
        ("Valor total", brl(report_total(rows, total_key))),
        ("Gerado em", datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%d/%m/%Y %H:%M")),
    ]
    if sector in ("orders", "finished"):
        summary.insert(2, ("Encomendas agrupadas", len({row["grupo"] for row in rows})))
    return {
        "title": title,
        "subtitle": "Relatório detalhado gerado a partir dos registros atuais do sistema.",
        "summary": summary,
        "sections": [{"title": sector_names[sector], "columns": columns, "rows": rows}],
    }


def tender_export_report(tender_id: int) -> dict:
    tender = tender_detail(tender_id)
    item_rows = [item_export_record(tender, item) for item in tender.get("items", [])]
    order_rows = [
        order_export_record(tender, item, order)
        for item in tender.get("items", [])
        for order in item.get("orders", [])
    ]
    item_by_id = {int(item.get("id") or 0): item.get("item") or "-" for item in tender.get("items", [])}
    document_rows = [{
        "tipo": attachment.get("tipo") or "-",
        "item": item_by_id.get(int(attachment.get("item_id") or 0), "Licitação"),
        "arquivo": attachment.get("filename") or "-",
        "content_type": attachment.get("content_type") or "-",
        "enviado_em": export_datetime(attachment.get("uploaded_at")),
    } for attachment in tender.get("attachments", [])]
    secret_count = sum(1 for item in tender.get("items", []) if item.get("valor_sigiloso"))
    proposal_date = export_datetime(tender.get("data_proposta"))
    summary = [
        ("Pregão", tender.get("pregão") or "-"),
        ("UASG", tender.get("uasg") or "-"),
        ("Órgão", tender.get("órgão") or "-"),
        ("Localidade", tender.get("localidade") or "-"),
        ("Plataforma", tender.get("plataforma") or "ComprasNet"),
        ("Modalidade", tender.get("modalidade") or "Pregão Eletrônico"),
        ("Data e hora", export_datetime(tender.get("data_limite")) or "-"),
        ("Data da proposta", proposal_date.date() if proposal_date else "-"),
        ("Status", tender.get("status") or "-"),
        ("Itens cadastrados", len(item_rows)),
        ("Itens com referência sigilosa", secret_count),
        ("Valor de referência", "Sigiloso" if secret_count else brl(tender.get("valor_total"))),
        ("Valor mínimo total", brl(tender.get("valor_minimo_total"))),
        ("Observação", tender.get("observação") or "-"),
        ("Gerado em", datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%d/%m/%Y %H:%M")),
    ]
    return {
        "title": f"LicitaUM - Pregão {tender.get('pregão')}",
        "subtitle": f"Relatório completo da licitação - UASG {tender.get('uasg')}",
        "summary": summary,
        "sections": [
            {"title": "Itens", "columns": ITEM_EXPORT_COLUMNS, "rows": item_rows},
            {"title": "Encomendas", "columns": ORDER_EXPORT_COLUMNS, "rows": order_rows},
            {"title": "Documentos", "columns": DOCUMENT_EXPORT_COLUMNS, "rows": document_rows},
        ],
    }


def export_temporary_path(base_name: str, output_format: str) -> Path:
    EXPORT_TEMP.mkdir(parents=True, exist_ok=True)
    return EXPORT_TEMP / f"{safe_name(base_name)}-{uuid.uuid4().hex}.{output_format}"


def read_body(handler: BaseHTTPRequestHandler) -> dict:
    length = int(handler.headers.get("Content-Length", "0"))
    if not length:
        return {}
    raw = handler.rfile.read(length).decode("utf-8")
    return json.loads(raw or "{}")


def save_tender(data: dict) -> int:
    fields = ["pregão", "uasg", "órgão", "localidade", "plataforma", "data_limite", "data_proposta", "status", "modalidade", "observação"]
    values = {key: data.get(key, "") for key in fields}
    values["plataforma"] = values["plataforma"] or "ComprasNet"
    values["status"] = values["status"] or "Em cadastro de preços"
    values["modalidade"] = values["modalidade"] or "Pregão Eletrônico"
    with connect() as con:
        if data.get("id"):
            sets = ", ".join(f"{field} = ?" for field in fields)
            con.execute(f"UPDATE tenders SET {sets} WHERE id = ?", [values[field] for field in fields] + [data["id"]])
            return int(data["id"])
        cur = con.execute(
            f"INSERT INTO tenders ({', '.join(fields)}) VALUES ({', '.join('?' for _ in fields)})",
            [values[field] for field in fields],
        )
        return int(cur.lastrowid)


def list_users() -> list[dict]:
    with connect() as con:
        rows = con.execute(
            "SELECT id, nome, email, perfil, ativo, created_at FROM users ORDER BY ativo DESC, nome"
        ).fetchall()
        return [row_to_dict(row) for row in rows]


def save_user(data: dict) -> int:
    nome = (data.get("nome") or "").strip()
    email = (data.get("email") or "").strip().lower()
    perfil = (data.get("perfil") or "Usuário").strip()
    senha = data.get("senha") or ""
    if not nome or not email:
        raise ValueError("Nome e e-mail são obrigatórios")
    if not data.get("id") and not senha:
        raise ValueError("Senha inicial é obrigatória para novos usuários")
    fields = ["nome", "email", "perfil", "ativo"]
    values = {
        "nome": nome,
        "email": email,
        "perfil": perfil,
        "ativo": 1 if str(data.get("ativo", "1")).lower() in ("1", "true", "on", "sim") else 0,
    }
    with connect() as con:
        existing = con.execute(
            "SELECT id FROM users WHERE email = ? AND id != ?",
            (email, int(data.get("id") or 0)),
        ).fetchone()
        if existing:
            raise ValueError("Já existe um usuário cadastrado com este e-mail")
        if data.get("id"):
            sets = [f"{field} = ?" for field in fields]
            args = [values[field] for field in fields]
            if senha:
                sets.append("senha_hash = ?")
                args.append(hash_password(senha))
            con.execute(f"UPDATE users SET {', '.join(sets)} WHERE id = ?", args + [data["id"]])
            return int(data["id"])
        cur = con.execute(
            "INSERT INTO users (nome, email, perfil, ativo, senha_hash) VALUES (?, ?, ?, ?, ?)",
            [values["nome"], values["email"], values["perfil"], values["ativo"], hash_password(senha)],
        )
        return int(cur.lastrowid)


def authenticate_user(data: dict) -> dict:
    email = (data.get("email") or "").strip().lower()
    senha = data.get("senha") or ""
    if not email or not senha:
      raise ValueError("Informe login e senha")
    with connect() as con:
        row = con.execute(
            "SELECT id, nome, email, perfil, ativo, senha_hash FROM users WHERE email = ?",
            (email,),
        ).fetchone()
    if not row or not int(row["ativo"] or 0) or row["senha_hash"] != hash_password(senha):
        raise PermissionError("Login ou senha inválidos")
    user = row_to_dict(row)
    user.pop("senha_hash", None)
    return user


def save_item(data: dict) -> int:
    fields = [
        "tender_id",
        "item",
        "marca",
        "modelo",
        "referência",
        "qtd",
        "valor_unitário",
        "link_usa",
        "link_br",
        "link_referência",
        "valor_cadastro",
        "valor_mínimo",
        "valor_ganho",
        "valor_sigiloso",
        "lote",
        "opção_produto",
        "selecionado_cadastro",
        "responsável_preço",
        "status",
        "observação",
    ]
    values = {key: data.get(key, "") for key in fields}
    for key in ("qtd", "valor_unitário", "valor_cadastro", "valor_mínimo", "valor_ganho"):
        values[key] = float(money(values[key]))
    values["valor_sigiloso"] = 1 if str(values["valor_sigiloso"]).lower() in ("1", "true", "on", "sim") else 0
    values["selecionado_cadastro"] = 1 if str(values["selecionado_cadastro"]).lower() in ("1", "true", "on", "sim", "") else 0
    values["status"] = values["status"] or "Em cadastro de preços"
    with connect() as con:
        if data.get("id"):
            sets = ", ".join(f"{field} = ?" for field in fields)
            con.execute(f"UPDATE items SET {sets} WHERE id = ?", [values[field] for field in fields] + [data["id"]])
            return int(data["id"])
        cur = con.execute(
            f"INSERT INTO items ({', '.join(fields)}) VALUES ({', '.join('?' for _ in fields)})",
            [values[field] for field in fields],
        )
        return int(cur.lastrowid)


def save_order_in_connection(con: sqlite3.Connection, data: dict, group_id: str | None = None) -> int:
    fields = ["item_id", "qtd_empenhada", "prazo_entrega", "ordem_fornecimento", "endereço_entrega", "nota_empenho", "status", "pagamento_recebido", "observação", "group_id", "origem", "órgão_solicitante"]
    values = {key: data.get(key, "") for key in fields}
    quantity = money(values["qtd_empenhada"])
    if quantity <= 0:
        raise ValueError("A quantidade empenhada deve ser maior que zero.")
    values["qtd_empenhada"] = float(quantity)
    values["status"] = values["status"] or "Pendente"
    values["origem"] = values["origem"] or "Empenho"
    values["pagamento_recebido"] = 1 if str(values["pagamento_recebido"]).lower() in ("1", "true", "on", "sim") else 0
    order_id = int(data["id"]) if data.get("id") else None
    item = con.execute("SELECT qtd FROM items WHERE id = ?", (values["item_id"],)).fetchone()
    if not item:
        raise ValueError("Item não encontrado.")
    committed = con.execute(
        """
        SELECT COALESCE(SUM(qtd_empenhada), 0)
        FROM orders
        WHERE item_id = ?
          AND id != COALESCE(?, -1)
          AND COALESCE(origem, 'Empenho') != 'Carona'
        """,
        (values["item_id"], order_id),
    ).fetchone()[0]
    available = money(item["qtd"]) - money(committed)
    if values["origem"] != "Carona" and quantity > available:
        raise ValueError(f"Quantidade superior ao saldo disponível ({available}).")

    if order_id:
        existing = con.execute(
            "SELECT group_id, origem, órgão_solicitante FROM orders WHERE id = ?",
            (order_id,),
        ).fetchone()
        if not existing:
            raise ValueError("Encomenda não encontrada.")
        values["group_id"] = group_id or values["group_id"] or existing["group_id"] or f"order-{order_id}"
        values["origem"] = data.get("origem") or existing["origem"] or "Empenho"
        values["órgão_solicitante"] = data.get("órgão_solicitante") or existing["órgão_solicitante"] or ""
        sets = ", ".join(f"{field} = ?" for field in fields)
        con.execute(
            f"UPDATE orders SET {sets} WHERE id = ?",
            [values[field] for field in fields] + [order_id],
        )
        return order_id

    values["group_id"] = group_id or values["group_id"] or uuid.uuid4().hex
    cur = con.execute(
        f"INSERT INTO orders ({', '.join(fields)}) VALUES ({', '.join('?' for _ in fields)})",
        [values[field] for field in fields],
    )
    return int(cur.lastrowid)


def save_order(data: dict) -> int:
    with connect() as con:
        return save_order_in_connection(con, data)


def save_orders_bulk(data: dict) -> list[int]:
    entries = data.get("items") or []
    if not entries:
        raise ValueError("Selecione ao menos um item.")
    common = {
        key: data.get(key, "")
        for key in ("prazo_entrega", "ordem_fornecimento", "endereço_entrega", "nota_empenho", "status", "pagamento_recebido", "observação", "origem", "órgão_solicitante")
    }
    group_id = uuid.uuid4().hex
    ids = []
    with connect() as con:
        for entry in entries:
            ids.append(
                save_order_in_connection(
                    con,
                    {**common, "item_id": entry.get("item_id"), "qtd_empenhada": entry.get("qtd_empenhada")},
                    group_id,
                )
            )
    return ids


def bulk_update_items(data: dict) -> int:
    item_ids = [int(item_id) for item_id in data.get("item_ids", [])]
    if not item_ids:
        return 0
    sets, args = [], []
    if data.get("status"):
        sets.append("status = ?")
        args.append(data["status"])
    if data.get("valor_ganho") not in (None, ""):
        sets.append("valor_ganho = ?")
        args.append(float(money(data["valor_ganho"])))
    if data.get("valor_sigiloso") not in (None, ""):
        sets.append("valor_sigiloso = ?")
        args.append(1 if str(data["valor_sigiloso"]).lower() in ("1", "true", "on", "sim") else 0)
    if not sets:
        return 0
    placeholders = ", ".join("?" for _ in item_ids)
    with connect() as con:
        con.execute(f"UPDATE items SET {', '.join(sets)} WHERE id IN ({placeholders})", args + item_ids)
    return len(item_ids)


def save_attachment(data: dict) -> int:
    UPLOADS.mkdir(exist_ok=True)
    filename = Path(data.get("filename") or "arquivo").name
    content_type = data.get("content_type") or mimetypes.guess_type(filename)[0] or "application/octet-stream"
    raw_data = data.get("data", "")
    if "," in raw_data:
        raw_data = raw_data.split(",", 1)[1]
    content = base64.b64decode(raw_data)
    stored_name = f"{uuid.uuid4().hex}_{safe_name(filename)}"
    (UPLOADS / stored_name).write_bytes(content)
    with connect() as con:
        cur = con.execute(
            """
            INSERT INTO attachments (tender_id, item_id, tipo, filename, stored_name, content_type)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                data.get("tender_id") or None,
                data.get("item_id") or None,
                data.get("tipo") or "Documento",
                filename,
                stored_name,
                content_type,
            ),
        )
        return int(cur.lastrowid)


def save_file_attachment(tender_id: int, file_path: Path, tipo: str) -> int:
    UPLOADS.mkdir(exist_ok=True)
    filename = file_path.name
    content_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"
    stored_name = f"{uuid.uuid4().hex}_{safe_name(filename)}"
    (UPLOADS / stored_name).write_bytes(file_path.read_bytes())
    with connect() as con:
        existing = con.execute(
            "SELECT id, stored_name FROM attachments WHERE tender_id = ? AND item_id IS NULL AND tipo = ? AND filename = ?",
            (tender_id, tipo, filename),
        ).fetchall()
        for row in existing:
            (UPLOADS / row["stored_name"]).unlink(missing_ok=True)
            con.execute("DELETE FROM attachments WHERE id = ?", (row["id"],))
        cur = con.execute(
            """
            INSERT INTO attachments (tender_id, item_id, tipo, filename, stored_name, content_type)
            VALUES (?, NULL, ?, ?, ?, ?)
            """,
            (tender_id, tipo, filename, stored_name, content_type),
        )
        return int(cur.lastrowid)


def get_attachment(attachment_id: int) -> dict:
    with connect() as con:
        row = con.execute("SELECT * FROM attachments WHERE id = ?", (attachment_id,)).fetchone()
        if not row:
            raise KeyError("Anexo não encontrado")
        return row_to_dict(row)


def number_words(value: Decimal) -> str:
    value = money(value).quantize(Decimal("0.01"))
    inteiro = int(value)
    centavos = int((value - inteiro) * 100)

    unidades = ["", "um", "dois", "três", "quatro", "cinco", "seis", "sete", "oito", "nove"]
    especiais = {
        10: "dez", 11: "onze", 12: "doze", 13: "treze", 14: "quatorze", 15: "quinze",
        16: "dezesseis", 17: "dezessete", 18: "dezoito", 19: "dezenove",
    }
    dezenas = ["", "", "vinte", "trinta", "quarenta", "cinquenta", "sessenta", "setenta", "oitenta", "noventa"]
    centenas = ["", "cento", "duzentos", "trezentos", "quatrocentos", "quinhentos", "seiscentos", "setecentos", "oitocentos", "novecentos"]

    def below_thousand(n: int) -> str:
        if n == 0:
            return ""
        if n == 100:
            return "cem"
        parts = []
        c, rest = divmod(n, 100)
        if c:
            parts.append(centenas[c])
        if rest:
            if rest < 10:
                parts.append(unidades[rest])
            elif rest < 20:
                parts.append(especiais[rest])
            else:
                d, u = divmod(rest, 10)
                parts.append(dezenas[d] + (f" e {unidades[u]}" if u else ""))
        return " e ".join(parts)

    def integer_words(n: int) -> str:
        if n == 0:
            return "zero"
        groups = []
        milhões, resto = divmod(n, 1_000_000)
        milhares, centenas_resto = divmod(resto, 1_000)
        if milhões:
            groups.append("um milhão" if milhões == 1 else f"{below_thousand(milhões)} milhões")
        if milhares:
            groups.append("mil" if milhares == 1 else f"{below_thousand(milhares)} mil")
        if centenas_resto:
            groups.append(below_thousand(centenas_resto))
        return ", ".join(groups[:-1]) + (" e " if len(groups) > 1 else "") + groups[-1]

    texto = integer_words(inteiro) + (" real" if inteiro == 1 else " reais")
    if centavos:
        texto += " e " + integer_words(centavos) + (" centavo" if centavos == 1 else " centavos")
    return texto


def selected_items(tender: dict, item_ids: list[int] | None) -> list[dict]:
    if not item_ids:
        return tender["items"]
    wanted = {int(item_id) for item_id in item_ids}
    return [item for item in tender["items"] if int(item["id"]) in wanted]


def proposal_link(item: dict) -> str:
    return str(item.get("link_referência") or item.get("link_br") or item.get("link_usa") or "").strip()


def generate_proposal(tender_id: int, item_ids: list[int] | None = None) -> Path:
    tender = tender_detail(tender_id)
    items = selected_items(tender, item_ids)
    if not items:
        raise ValueError("Selecione pelo menos um item para gerar a proposta.")
    missing_links = [str(item.get("item") or item.get("modelo") or item["id"]) for item in items if not proposal_link(item)]
    if missing_links:
        raise ValueError("Informe o link do fornecedor antes de gerar a proposta. Itens sem link: " + ", ".join(missing_links))
    OUTPUTS.mkdir(exist_ok=True)
    if not TEMPLATE_COPY.exists():
        raise FileNotFoundError("Modelo de proposta não encontrado.")

    doc = Document(TEMPLATE_COPY)
    set_paragraph_text(
        doc.paragraphs[1],
        f"Ao Órgão UASG {tender['uasg']} - {tender['órgão']}. Apresentamos nossa proposta de preços.",
    )

    table = doc.tables[0]
    while len(table.rows) > 2:
        table._tbl.remove(table.rows[-1]._tr)
    if len(table.rows) < 2:
        table.add_row()
    total = Decimal("0")
    for idx, item in enumerate(items):
        row = table.rows[1].cells if idx == 0 else table.add_row().cells
        total_item = money(item["qtd"]) * effective_unit_value(item)
        total += total_item
        set_cell_text(row[0], str(item.get("item") or ""))
        set_cell_text(row[1], str(item.get("marca") or ""))
        set_cell_text(row[2], str(item.get("modelo") or item.get("referência") or ""))
        set_cell_text(row[3], proposal_link(item))
        set_cell_text(row[4], str(item.get("qtd") or ""))
        set_cell_text(row[5], brl(effective_unit_value(item)))
        set_cell_text(row[6], brl(total_item))

    for paragraph in doc.paragraphs:
        text = paragraph.text
        if "Valor total da proposta:" in text:
            set_paragraph_text(paragraph, re.sub(r"Valor total da proposta:.*", f"Valor total da proposta: {brl(total)}", text))
        elif "O valor total dessa proposta é de" in text:
            set_paragraph_text(paragraph, f"O valor total dessa proposta é de {brl(total)} ({number_words(total)}).")

    when_date = current_business_date()
    months = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]
    date_line = f"Muqui/ES, {when_date.day} de {months[when_date.month - 1].capitalize()} de {when_date.year}"
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Muqui/ES,"):
            set_paragraph_text(paragraph, date_line)
            break

    filename = proposal_filename(tender, items)
    out = OUTPUTS / filename
    doc.save(out)
    return out


def find_soffice() -> str | None:
    for candidate in (
        shutil.which("soffice"),
        shutil.which("libreoffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ):
        if candidate and Path(candidate).exists():
            return str(candidate)
    return None


def convert_docx_to_pdf(docx_path: Path) -> Path:
    pdf_path = docx_path.with_suffix(".pdf")
    soffice = find_soffice()
    if soffice:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
            check=True,
            capture_output=True,
            text=True,
            timeout=60,
        )
        if pdf_path.exists():
            return pdf_path

    winword_exists = any(
        Path(candidate).exists()
        for candidate in (
            r"C:\Program Files\Microsoft Office\root\Office16\WINWORD.EXE",
            r"C:\Program Files (x86)\Microsoft Office\root\Office16\WINWORD.EXE",
        )
    )
    if not winword_exists:
        raise RuntimeError("Não foi possível converter para PDF. Instale LibreOffice ou Microsoft Word neste computador para habilitar a conversão automática.")

    ps = f"""
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $doc = $word.Documents.Open('{str(docx_path).replace("'", "''")}', $false, $true)
    $doc.ExportAsFixedFormat('{str(pdf_path).replace("'", "''")}', 17)
    $doc.Close(0)
    $word.Quit()
    """
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True, capture_output=True, text=True, timeout=20)
    except Exception as exc:
        raise RuntimeError("Não foi possível converter para PDF. Instale LibreOffice ou Microsoft Word neste computador para habilitar a conversão automática.") from exc
    if not pdf_path.exists():
        raise RuntimeError("A conversão para PDF não retornou um arquivo.")
    return pdf_path


class App(BaseHTTPRequestHandler):
    def send_json(self, data, code: int = 200) -> None:
        payload = json.dumps(data, ensure_ascii=False, default=str).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def send_file(self, path: Path, content_type: str = "application/octet-stream", download_name: str | None = None, cache_control: str | None = None) -> None:
        payload = path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        if cache_control:
            self.send_header("Cache-Control", cache_control)
        if download_name:
            self.send_header("Content-Disposition", f'attachment; filename="{download_name}"')
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def send_export(self, report: dict, output_format: str, base_name: str) -> None:
        output = export_temporary_path(base_name, output_format)
        try:
            if output_format == "xlsx":
                from report_exports import generate_xlsx
                generate_xlsx(report, output)
                content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            elif output_format == "pdf":
                from report_exports import generate_pdf
                generate_pdf(report, output)
                content_type = "application/pdf"
            else:
                raise ValueError("Formato de exportação inválido.")
            self.send_file(output, content_type, f"{safe_name(base_name)}.{output_format}")
        finally:
            output.unlink(missing_ok=True)

    def do_GET(self) -> None:
        try:
            parsed = urlparse(self.path)
            path = parsed.path
            qs = parse_qs(parsed.query)
            if path == "/":
                return self.send_file(STATIC / "index.html", "text/html; charset=utf-8", cache_control="no-store")
            if path.startswith("/static/"):
                file_path = STATIC / path.removeprefix("/static/")
                types = {
                    ".css": "text/css; charset=utf-8",
                    ".js": "application/javascript; charset=utf-8",
                    ".svg": "image/svg+xml",
                    ".png": "image/png",
                    ".jpg": "image/jpeg",
                    ".jpeg": "image/jpeg",
                    ".webp": "image/webp",
                }
                return self.send_file(file_path, types.get(file_path.suffix, "application/octet-stream"), cache_control="no-store")
            if path == "/api/meta":
                return self.send_json({"status": STATUS, "current_date": current_business_date().isoformat()})
            if path == "/api/users":
                return self.send_json(list_users())
            if path == "/api/attachments":
                tender_id = int(qs.get("tender_id", ["0"])[0] or 0)
                item_id = int(qs.get("item_id", ["0"])[0] or 0)
                return self.send_json(attachments_for(tender_id or None, item_id or None))
            if path == "/api/state":
                return self.send_json(application_state())
            if path == "/api/tenders":
                return self.send_json(all_tenders(qs.get("q", [""])[0], qs.get("status", [""])[0]))
            if path.startswith("/api/tenders/"):
                tender_id = int(path.split("/")[-1])
                return self.send_json(tender_detail(tender_id))
            if path.startswith("/attachments/"):
                attachment_id = int(path.strip("/").split("/")[1])
                attachment = get_attachment(attachment_id)
                return self.send_file(UPLOADS / attachment["stored_name"], attachment.get("content_type") or "application/octet-stream", attachment["filename"])
            if path.startswith("/export/tender/"):
                tender_id = int(path.strip("/").split("/")[-1])
                output_format = qs.get("format", ["xlsx"])[0].lower()
                report = tender_export_report(tender_id)
                tender = tender_detail(tender_id)
                base_name = f"Relatorio Pregao {tender.get('pregão')} UASG {tender.get('uasg')}"
                return self.send_export(report, output_format, base_name)
            if path.startswith("/export/sector/"):
                sector = path.strip("/").split("/")[-1]
                output_format = qs.get("format", ["xlsx"])[0].lower()
                raw_ids = qs.get("ids", [None])[0]
                selected_ids = {int(value) for value in raw_ids.split(",") if value} if raw_ids is not None else None
                report = sector_export_report(sector, selected_ids)
                sector_label = {"won": "Itens Ganhos", "proposal": "Itens em Proposta", "orders": "Encomendas", "finished": "Finalizados"}.get(sector, "Relatorio")
                base_name = f"LicitaUM {sector_label} {current_business_date().isoformat()}"
                return self.send_export(report, output_format, base_name)
            if path.startswith("/proposal/"):
                tender_id = int(path.split("/")[-1])
                item_ids = [int(item_id) for item_id in qs.get("items", [""])[0].split(",") if item_id]
                out = generate_proposal(tender_id, item_ids)
                if qs.get("format", ["docx"])[0].lower() == "pdf":
                    pdf = convert_docx_to_pdf(out)
                    save_file_attachment(tender_id, pdf, "Proposta")
                    return self.send_file(pdf, "application/pdf", pdf.name)
                save_file_attachment(tender_id, out, "Proposta")
                return self.send_file(out, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", out.name)
            self.send_error(404)
        except Exception as exc:
            self.send_json({"error": str(exc)}, 500)

    def do_POST(self) -> None:
        try:
            path = urlparse(self.path).path
            parsed = urlparse(self.path)
            qs = parse_qs(parsed.query)
            data = read_body(self)
            if path.startswith("/export/sector/"):
                sector = path.strip("/").split("/")[-1]
                output_format = qs.get("format", ["xlsx"])[0].lower()
                selected_ids = {int(value) for value in data.get("ids", [])}
                report = sector_export_report(sector, selected_ids)
                sector_label = {"won": "Itens Ganhos", "proposal": "Itens em Proposta", "orders": "Encomendas", "finished": "Finalizados"}.get(sector, "Relatorio")
                base_name = f"LicitaUM {sector_label} {current_business_date().isoformat()}"
                return self.send_export(report, output_format, base_name)
            if path == "/api/login":
                return self.send_json(authenticate_user(data))
            if path == "/api/tenders":
                return self.send_json({"id": save_tender(data)})
            if path == "/api/users":
                return self.send_json({"id": save_user(data)})
            if path == "/api/items":
                return self.send_json({"id": save_item(data)})
            if path == "/api/orders":
                return self.send_json({"id": save_order(data)})
            if path == "/api/orders/bulk":
                return self.send_json({"ids": save_orders_bulk(data)})
            if path == "/api/items/bulk":
                return self.send_json({"updated": bulk_update_items(data)})
            if path == "/api/attachments":
                return self.send_json({"id": save_attachment(data)})
            self.send_error(404)
        except Exception as exc:
            self.send_json({"error": str(exc)}, 500)

    def do_DELETE(self) -> None:
        try:
            path = urlparse(self.path).path
            parts = path.strip("/").split("/")
            if len(parts) == 3 and parts[0] == "api":
                table = {"tenders": "tenders", "items": "items", "orders": "orders", "attachments": "attachments", "users": "users"}.get(parts[1])
                if table:
                    with connect() as con:
                        if table == "attachments":
                            row = con.execute("SELECT stored_name FROM attachments WHERE id = ?", (int(parts[2]),)).fetchone()
                            if row:
                                (UPLOADS / row["stored_name"]).unlink(missing_ok=True)
                        if table == "tenders":
                            rows = con.execute("SELECT stored_name FROM attachments WHERE tender_id = ?", (int(parts[2]),)).fetchall()
                            for row in rows:
                                (UPLOADS / row["stored_name"]).unlink(missing_ok=True)
                        if table == "items":
                            rows = con.execute("SELECT stored_name FROM attachments WHERE item_id = ?", (int(parts[2]),)).fetchall()
                            for row in rows:
                                (UPLOADS / row["stored_name"]).unlink(missing_ok=True)
                        con.execute(f"DELETE FROM {table} WHERE id = ?", (int(parts[2]),))
                    return self.send_json({"ok": True})
            self.send_error(404)
        except Exception as exc:
            self.send_json({"error": str(exc)}, 500)


class AppServer(ThreadingHTTPServer):
    daemon_threads = True
    allow_reuse_address = True


if __name__ == "__main__":
    init_db()
    ensure_schema()
    host = os.environ.get("LICITAUM_HOST", "127.0.0.1")
    port = int(os.environ.get("LICITAUM_PORT", "8765"))
    server = AppServer((host, port), App)
    print(f"LicitaUM aberto em http://{host}:{port}")
    server.serve_forever()
